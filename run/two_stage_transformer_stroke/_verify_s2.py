from openpyxl import load_workbook
from pathlib import Path
import time
p=Path(r"G:/02block_result/33_AKI/two_stage_transformer_40041421/summary_results/Tables/Table S2-MIMIC. Features used in Transformer.xlsx")
wb=load_workbook(p, read_only=True)
ws=wb.active
rows=list(ws.iter_rows(values_only=True))
print("title:", rows[0][0])
print("header:", rows[1])
statics=[r for r in rows[2:] if r and r[0]=="Static (baseline broadcast)"]
dyn=[r for r in rows[2:] if r and r[0]=="Dynamic (hourly)"]
print("n_dynamic_rows:", len(dyn))
print("n_static_rows:", len(statics))
print("static:", [r[1] for r in statics])
bmi=[r for r in rows if r and r[1] and "BMI" in str(r[1]).upper()]
alb=[r for r in rows if r and r[1] and "Albumin" in str(r[1])]
print("BMI rows:", bmi)
print("Albumin rows:", alb)
print("mtime:", time.strftime("%Y-%m-%d %H:%M:%S", time.localtime(p.stat().st_mtime)), "size:", p.stat().st_size)
