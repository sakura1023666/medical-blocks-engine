#!/usr/bin/env python3
"""Rebuild Table S2: dynamic coverage audit + final 18 static features."""
from __future__ import annotations

import csv
from pathlib import Path

import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter

def _proj() -> Path:
    for cand in (
        Path("G:/02block_result/33_AKI/two_stage_transformer_40041421"),
        Path("/mnt/g/02block_result/33_AKI/two_stage_transformer_40041421"),
    ):
        if cand.exists():
            return cand
    raise FileNotFoundError("AKI TST project root not found")


PROJ = _proj()
audit = PROJ / "_shared/step05_tst_timeseries/_tst_feature_coverage_audit.csv"
_npz_cands = list((PROJ / "by_unit").glob("*/step09_tst_train_eval/Tables/npz/test.npz"))
_npz_cands = [p for p in _npz_cands if "L120_B_twostage" in str(p)]
if not _npz_cands:
    raise FileNotFoundError("L120_B_twostage test.npz not found")
npz = sorted(_npz_cands, key=lambda p: ("success" not in str(p).lower(), len(str(p))))[0]
out = PROJ / "summary_results/Tables/Table S2-MIMIC. Features used in Transformer.xlsx"
print("using npz", npz)


def strip_us(s: object) -> str:
    t = "" if s is None else str(s)
    if t.lower().startswith("static_"):
        t = t[7:]
    return t.replace("_", " ").replace("  ", " ").strip()


names = list(np.load(npz, allow_pickle=True)["feature_names"])
static = [n for n in names if str(n).startswith("static_")]
dynamic_in_model = [n for n in names if not str(n).startswith("static_")]
assert len(names) == 62 and len(static) == 18 and len(dynamic_in_model) == 44, (
    len(names),
    len(static),
    len(dynamic_in_model),
)

with audit.open(encoding="utf-8", newline="") as f:
    rows_audit = list(csv.DictReader(f))

combined = []
for r in rows_audit:
    feat = r.get("feature") or r.get("Feature") or ""
    miss = r.get("day1_missing_pct") or r.get("day1 missing pct") or ""
    thr = r.get("threshold_pct") or r.get("threshold pct") or ""
    act = r.get("action") or r.get("Action") or ""
    combined.append(
        {
            "Section": "Dynamic (hourly)",
            "Feature": strip_us(feat),
            "In final model": "Yes" if act in ("KEPT", "KEPT_FORCE") else "No",
            "Day1 missing %": miss,
            "Threshold %": thr,
            "Audit action": act,
        }
    )

for s in static:
    combined.append(
        {
            "Section": "Static (baseline broadcast)",
            "Feature": strip_us(s),
            "In final model": "Yes",
            "Day1 missing %": "—",
            "Threshold %": "—",
            "Audit action": "STATIC",
        }
    )

title = "Table S2. Features used in Transformer (MIMIC AKI)"
footnotes = [
    f"Final model input F={len(names)} = {len(dynamic_in_model)} dynamic + {len(static)} static features (from L120 test.npz feature_names).",
    "Dynamic block: day-1 coverage audit after dense 24h grid and forward-fill; KEPT/KEPT_FORCE enter the model; DROPPED do not.",
    "Static block: admission/baseline variables broadcast across hours; BMI and Albumin were listed in config candidates but were NOT present in the final 62-feature tensor and are omitted here.",
    "Analysis cohort N = 14,450 after day1–5 patient missing-data filter (aligned with Table 1 / Figure 1).",
    "Score comparators (APSIII/SOFA/SAPSII/OASIS) are excluded from model inputs to avoid leakage vs Figure 2 / Table 2.",
]

FONT = "Times New Roman"
thick = Side(style="medium", color="000000")
thin = Side(style="thin", color="000000")
border_header = Border(top=thick, bottom=thin)
border_last = Border(bottom=thick)
border_foot_top = Border(top=thin)
font_title = Font(name=FONT, size=12, bold=True)
font_hdr = Font(name=FONT, size=12, bold=True)
font_body = Font(name=FONT, size=12)
align_c = Alignment(horizontal="center", vertical="center", wrap_text=True)
align_l = Alignment(horizontal="left", vertical="center", wrap_text=True)

cols = ["Section", "Feature", "In final model", "Day1 missing %", "Threshold %", "Audit action"]
wb = Workbook()
ws = wb.active
ws.title = "Table"
ws.sheet_view.showGridLines = False
nc = len(cols)
ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=nc)
c = ws.cell(1, 1, title)
c.font = font_title
c.alignment = align_c
for j, col in enumerate(cols, 1):
    cell = ws.cell(2, j, col)
    cell.font = font_hdr
    cell.alignment = align_c
    cell.border = border_header
for i, row in enumerate(combined):
    excel_r = 3 + i
    is_last = i == len(combined) - 1
    for j, col in enumerate(cols, 1):
        cell = ws.cell(excel_r, j, row[col])
        cell.font = font_body
        cell.alignment = align_l if j <= 2 else align_c
        if is_last:
            cell.border = border_last
last_data = 2 + len(combined)
for k, note in enumerate(footnotes):
    fr = last_data + 1 + k
    ws.merge_cells(start_row=fr, start_column=1, end_row=fr, end_column=nc)
    cell = ws.cell(fr, 1, note)
    cell.font = font_body
    cell.alignment = align_l
    if k == 0:
        cell.border = border_foot_top
for j, col in enumerate(cols, 1):
    maxlen = max(len(col), max(len(str(r[col])) for r in combined))
    ws.column_dimensions[get_column_letter(j)].width = min(max(maxlen + 2, 10), 42)
out.parent.mkdir(parents=True, exist_ok=True)
wb.save(out)
print("Wrote", out)
print("n_rows", len(combined))
print("static", [strip_us(s) for s in static])
