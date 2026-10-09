#!/usr/bin/env python3
"""Remove NHANES survey weight / design rows from published xlsx (column A)."""
import os
import re
import sys

STUDY = "/mnt/g/02block_result/24_diabetes/incidence_38341157"
WT_ROW = re.compile(r"^(WT[A-Z0-9_]+|SDMV[A-Z0-9_]*|Source_File|new_Weight|new_weight)$", re.I)


def strip_file(path):
    try:
        import openpyxl
    except ImportError:
        print("openpyxl required", file=sys.stderr)
        return False
    try:
        wb = openpyxl.load_workbook(path)
    except Exception as e:
        print("skip corrupt:", path, e, file=sys.stderr)
        return False
    ws = wb.active
    drop = []
    for i in range(1, ws.max_row + 1):
        v = ws.cell(row=i, column=1).value
        if v is not None and WT_ROW.match(str(v).strip()):
            drop.append(i)
    if not drop:
        return False
    for r in sorted(drop, reverse=True):
        ws.delete_rows(r, 1)
    wb.save(path)
    return True


def main():
    n = 0
    bi = os.path.join(STUDY, "by_index")
    for idx in os.listdir(bi):
        for sub in ("Tables", os.path.join("NHANES", "Tables"), "NHANES"):
            td = os.path.join(bi, idx, sub)
            if not os.path.isdir(td):
                continue
            for fn in os.listdir(td):
                if not fn.endswith(".xlsx"):
                    continue
                if "NHANES" not in fn and "nhanes" not in fn.lower():
                    continue
                p = os.path.join(td, fn)
                if strip_file(p):
                    n += 1
                    print("stripped:", p)
    print("total stripped:", n)


if __name__ == "__main__":
    main()
