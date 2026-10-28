#!/usr/bin/env python3
"""Nature statistics + figure delivery scan for a published project root.

Machine half of Layer D (nature-statistics / nature-figure). Repairs raster
DPI below 600 from the matching PDF. Does not invent statistical conclusions.
Agent still fills narrative QC and the midterm content.json.

    python3 python/nature_two_aspect_scan.py --project <root> [--repair]
"""
from __future__ import annotations

import argparse
import glob
import os
import subprocess
import sys
from datetime import date


def _roots(project: str) -> list[str]:
    cands = [
        os.path.join(project, "summary_results"),
        os.path.join(project, "summary_result"),
        project,
    ]
    out = []
    for c in cands:
        if os.path.isdir(os.path.join(c, "Figures", "pdf")) or os.path.isdir(os.path.join(c, "Tables")):
            out.append(c)
            break
    return out


def _dpi(path: str):
    from PIL import Image

    im = Image.open(path)
    dpi = im.info.get("dpi")
    im.close()
    if isinstance(dpi, tuple) and dpi:
        return min(float(x) for x in dpi)
    return None


def scan_figures(pub: str) -> dict:
    fig = os.path.join(pub, "Figures")
    pdfs = sorted(glob.glob(os.path.join(fig, "pdf", "Figure*.pdf")))
    stems = [os.path.splitext(os.path.basename(p))[0] for p in pdfs]
    bad_dpi = []
    missing = []
    tif = glob.glob(os.path.join(fig, "**", "*.tif"), recursive=True)
    for st in stems:
        for ext in ("png", "tiff"):
            p = os.path.join(fig, ext, st + "." + ext)
            if not os.path.isfile(p) or os.path.getsize(p) < 800:
                missing.append(os.path.basename(p))
                continue
            d = _dpi(p)
            if d is None or d < 599.0:
                bad_dpi.append((p, d))
    info = os.path.join(fig, "image_information")
    ghosts = []
    if os.path.isdir(info):
        for md in glob.glob(os.path.join(info, "Figure*.md")):
            st = os.path.splitext(os.path.basename(md))[0]
            if not os.path.isfile(os.path.join(fig, "pdf", st + ".pdf")):
                ghosts.append(os.path.basename(md))
    return {"stems": stems, "bad_dpi": bad_dpi, "missing": missing, "tif": tif, "ghosts": ghosts}


def repair(pub: str, bad_dpi: list, dpi: int, engine: str) -> list[str]:
    script = os.path.join(engine, "python", "pub_figure_rasterize.py")
    done = []
    stems = []
    for path, _d in bad_dpi:
        st = os.path.splitext(os.path.basename(path))[0]
        if st not in stems:
            stems.append(st)
    fig = os.path.join(pub, "Figures")
    for st in stems:
        pdf = os.path.join(fig, "pdf", st + ".pdf")
        if not os.path.isfile(pdf):
            continue
        cmd = [
            sys.executable, script, "--pdf", pdf, "--dpi", str(dpi),
            "--png", os.path.join(fig, "png", st + ".png"),
            "--tiff", os.path.join(fig, "tiff", st + ".tiff"),
        ]
        subprocess.check_call(cmd)
        done.append(st)
    return done


def scan_xlsx(pub: str) -> list[str]:
    from openpyxl import load_workbook

    bad = []
    for p in sorted(glob.glob(os.path.join(pub, "Tables", "*.xlsx"))):
        try:
            wb = load_workbook(p)
            wb.close()
        except Exception as exc:  # noqa: BLE001
            bad.append(f"{os.path.basename(p)}: {type(exc).__name__}")
    return bad


def repair_xlsx(pub: str, engine: str) -> list[str]:
    """Run engine xlsx_strip on Tables that openpyxl cannot open."""
    strip = os.path.join(engine, "R", "pub_xlsx_tools", "xlsx_strip.py")
    if not os.path.isfile(strip):
        return []
    fixed = []
    for p in sorted(glob.glob(os.path.join(pub, "Tables", "*.xlsx"))):
        try:
            from openpyxl import load_workbook

            wb = load_workbook(p)
            wb.close()
            continue
        except Exception:  # noqa: BLE001
            pass
        try:
            subprocess.check_call([sys.executable, strip, p])
            from openpyxl import load_workbook

            wb = load_workbook(p)
            wb.close()
            fixed.append(os.path.basename(p))
        except Exception as exc:  # noqa: BLE001
            fixed.append(f"{os.path.basename(p)} (still broken: {type(exc).__name__})")
    return fixed


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--project", required=True)
    ap.add_argument("--repair", action="store_true")
    ap.add_argument("--dpi", type=int, default=600)
    args = ap.parse_args()
    project = os.path.abspath(args.project)
    pubs = _roots(project)
    if not pubs:
        print("no published Figures/Tables", file=sys.stderr)
        return 2
    engine = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    day = date.today().isoformat()
    lines = []
    n_p0 = 0
    for pub in pubs:
        fig = scan_figures(pub)
        repaired = []
        if args.repair and fig["bad_dpi"]:
            repaired = repair(pub, fig["bad_dpi"], args.dpi, engine)
            fig = scan_figures(pub)
        xfixed = []
        if args.repair:
            xfixed = repair_xlsx(pub, engine)
        xbad = scan_xlsx(pub)
        p0 = []
        if fig["bad_dpi"]:
            p0.append("DPI<600: " + "; ".join(os.path.basename(p) for p, _ in fig["bad_dpi"]))
        if fig["missing"]:
            p0.append("缺 png/tiff: " + "; ".join(fig["missing"][:8]))
        if fig["ghosts"]:
            p0.append("image_information 无对应 pdf: " + "; ".join(fig["ghosts"]))
        if xbad:
            p0.append("xlsx 打不开: " + "; ".join(xbad[:8]))
        n_p0 += len(p0)
        lines.append(f"## `{pub}`")
        lines.append(f"- 图 {len(fig['stems'])} 张；DPI 不合格 {len(fig['bad_dpi'])}；本次重导 {len(repaired)}")
        if repaired:
            lines.append("- 已从 PDF 重导 600 dpi：" + "; ".join(repaired))
        if xfixed:
            lines.append("- 已 xlsx_strip 修复：" + "; ".join(xfixed))
        lines.append(f"- .tif {len(fig['tif'])}；幽灵说明 {len(fig['ghosts'])}；xlsx 打不开 {len(xbad)}")
        if p0:
            lines.append("- 机器 P0：" + " | ".join(p0))
        else:
            lines.append("- 机器 P0：无（横断面 Incidence、图号错置、中介与主表口径仍须人工过清单）")
        lines.append("")
    reports = os.path.join(project, "reports")
    os.makedirs(reports, exist_ok=True)
    body = "\n".join([
        "# Nature 两方面机器扫描",
        "",
        f"- 项目：`{project}`",
        f"- 日期：{day}",
        "- 范围：nature-figure（DPI、四目录、.tiff、幽灵说明）与 nature-statistics（xlsx 能否打开）。",
        "- 叙述性 P0 不在本脚本内下结论。",
        "",
        *lines,
    ])
    out = os.path.join(reports, f"nature_two_aspect_scan_{day}.md")
    with open(out, "w", encoding="utf-8") as fh:
        fh.write(body)
    print(out)
    print("machine_p0", n_p0)
    return 1 if n_p0 else 0


if __name__ == "__main__":
    raise SystemExit(main())
