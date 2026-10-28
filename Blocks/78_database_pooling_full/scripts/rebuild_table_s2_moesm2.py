#!/usr/bin/env python3
"""Rebuild Table S2 to match MOESM2 layout (equal-n four-DB training + 95% CIs).

Columns: Prediction | Validation dataset | LSTM | Transformer | TPC
Rows: MAPE then AUROC × AmsterdamUMC, eICU, HiRID, MIMIC-IV
Best model per row in bold (MAPE: min; AUROC: max).
"""
from __future__ import annotations

import argparse
import sys
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter

DB_ORDER = ["amsterdam", "eicu", "hirid", "mimic4"]  # MOESM2 order
DB_LAB = {"amsterdam": "AmsterdamUMC", "mimic4": "MIMIC-IV", "eicu": "eICU", "hirid": "HiRID"}
MODELS = ("LSTM", "Transformer", "TPC")


def _fmt_ci_en(est, lo, hi, digits: int) -> str:
    if est is None or (isinstance(est, float) and not np.isfinite(est)):
        return "–"
    if lo is None or hi is None or not (np.isfinite(lo) and np.isfinite(hi)):
        return f"{float(est):.{digits}f}"
    return f"{float(est):.{digits}f} [{float(lo):.{digits}f}–{float(hi):.{digits}f}]"


def _fmt_n(n) -> str:
    try:
        return f"{int(n):,}"
    except Exception:
        return ""


def _last_test(mpath: Path) -> dict:
    df = pd.read_csv(mpath)
    te = df[df["set"].astype(str).str.lower() == "test"].copy()
    te = te[te["epoch"] == te["epoch"].max()]
    out = {"n_train": int(df["n_train"].iloc[0])}
    for _, r in te.iterrows():
        task = str(r["task"]).lower()
        db = str(r["source_dataset"])
        if "mort" in task:
            out.setdefault(db, {})["auc"] = float(r["auc"]) if pd.notna(r.get("auc")) else float("nan")
        if "los" in task:
            out.setdefault(db, {})["mape"] = float(r["mape"]) if pd.notna(r.get("mape")) else float("nan")
    return out


def _load_metrics_grid(equaln_root: Path) -> tuple[dict, int | None]:
    grid = {m: {} for m in MODELS}
    n_train = None
    for mpath in sorted(equaln_root.glob("results/**/metrics_multitask.csv")):
        model = "TPC"
        for part in mpath.parts:
            pl = part.lower()
            if pl == "lstm":
                model = "LSTM"
            elif pl == "transformer":
                model = "Transformer"
            elif pl == "tpc":
                model = "TPC"
        spec = _last_test(mpath)
        n_train = spec.get("n_train", n_train)
        for db in DB_ORDER:
            if db in spec:
                grid[model][db] = spec[db]
    return grid, n_train


def _load_ci(ci_csv: Path) -> pd.DataFrame:
    if not ci_csv.exists():
        return pd.DataFrame()
    df = pd.read_csv(ci_csv)
    df["model"] = df["model"].astype(str).str.lower().map(
        {"lstm": "LSTM", "transformer": "Transformer", "tpc": "TPC"}
    )
    return df


def build_s2(equaln_root: Path, ci_csv: Path) -> tuple[pd.DataFrame, pd.DataFrame, int | None]:
    """Return (display_df, numeric_best_mask_source, n_train).

    display_df columns: Prediction, Validation dataset, LSTM, Transformer, TPC
    best_num: same shape numeric for bold (MAPE or AUROC point).
    """
    grid, n_train = _load_metrics_grid(equaln_root)
    ci = _load_ci(ci_csv)
    rows = []
    best_rows = []
    for task, key, digits, prefer in (
        ("Remaining length of stay, MAPE [95% CI]", "mape", 1, "min"),
        ("Mortality, AUROC [95% CI]", "auc", 3, "max"),
    ):
        for db in DB_ORDER:
            row = {"Prediction": task, "Validation dataset": DB_LAB[db]}
            nums = {}
            for model in MODELS:
                # Prefer stay-level point+CI (same as Table 2/3); fall back to metrics CSV.
                est = grid.get(model, {}).get(db, {}).get(key, float("nan"))
                lo = hi = float("nan")
                if len(ci):
                    sub = ci[(ci["model"] == model) & (ci["eval_on"] == db)]
                    if len(sub):
                        r0 = sub.iloc[0]
                        if key == "auc":
                            est = float(r0.get("auroc_stay", est))
                            lo = float(r0.get("auroc_stay_lo_delong", np.nan))
                            hi = float(r0.get("auroc_stay_hi_delong", np.nan))
                        else:
                            est = float(r0.get("mape_stay_mean", est))
                            lo = float(r0.get("mape_lo_boot10k", np.nan))
                            hi = float(r0.get("mape_hi_boot10k", np.nan))
                row[model] = _fmt_ci_en(est, lo, hi, digits)
                nums[model] = float(est) if est is not None and np.isfinite(est) else float("nan")
            rows.append(row)
            best_rows.append({"prefer": prefer, **nums})
    return pd.DataFrame(rows), pd.DataFrame(best_rows), n_train


def write_s2_xlsx(display: pd.DataFrame, best_src: pd.DataFrame, out: Path, title: str, footnote: str):
    wb = Workbook()
    ws = wb.active
    ws.title = "Table"
    ws.sheet_view.showGridLines = False
    nc = len(display.columns)
    thick = Side(style="medium", color="000000")
    thin = Side(style="thin", color="000000")
    font_title = Font(name="Times New Roman", bold=True, size=12)
    font_h = Font(name="Times New Roman", bold=True, size=12)
    font_b = Font(name="Times New Roman", size=12)
    font_bold = Font(name="Times New Roman", bold=True, size=12)
    font_fn = Font(name="Times New Roman", size=10)
    align_c = Alignment(horizontal="center", vertical="center", wrap_text=True)
    align_l = Alignment(horizontal="left", vertical="center", wrap_text=True)

    cell = ws.cell(1, 1, title)
    cell.font = font_title
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=nc)
    ws.row_dimensions[1].height = 32

    for j, col in enumerate(display.columns, 1):
        cell = ws.cell(2, j, col)
        cell.font = font_h
        cell.alignment = align_c
        cell.border = Border(top=thick, bottom=thin)

    n = len(display)
    for i, row in enumerate(display.itertuples(index=False), 3):
        last = i == 2 + n
        prefer = best_src.iloc[i - 3]["prefer"]
        vals_num = [best_src.iloc[i - 3][m] for m in MODELS]
        if prefer == "min":
            best_i = int(np.nanargmin(vals_num)) if np.any(np.isfinite(vals_num)) else -1
        else:
            best_i = int(np.nanargmax(vals_num)) if np.any(np.isfinite(vals_num)) else -1
        for j, val in enumerate(row, 1):
            cell = ws.cell(i, j, val)
            is_model_col = j >= 3
            model_idx = j - 3
            bold = is_model_col and model_idx == best_i
            cell.font = font_bold if bold else font_b
            cell.alignment = align_l if j == 1 else align_c
            if last:
                cell.border = Border(bottom=thick)

    fr = 2 + n + 2
    cell = ws.cell(fr, 1, footnote)
    cell.font = font_fn
    cell.alignment = align_l
    ws.merge_cells(start_row=fr, start_column=1, end_row=fr, end_column=nc)
    cell.border = Border(top=thin)

    widths = [42, 18, 28, 28, 28]
    for j, w in enumerate(widths[:nc], 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--extra", required=True, type=Path, help="Output_paper_extras_physclip")
    ap.add_argument("--out", required=True, type=Path)
    ap.add_argument("--cohort-label", default="ACG cohort")
    args = ap.parse_args()

    equaln = args.extra / "liver_model_benchmark_equaln"
    ci_csv = args.extra / "ci_liver_model_benchmark_equaln" / "Table_CI_DeLong_bootstrap10k.csv"
    if not equaln.exists():
        raise SystemExit(f"missing {equaln}")

    display, best_src, n_train = build_s2(equaln, ci_csv)
    if display.empty:
        raise SystemExit("Table S2 empty — no equaln metrics")

    n_models_ci = 0
    if ci_csv.exists():
        n_models_ci = pd.read_csv(ci_csv)["model"].nunique()

    title = (
        "Table S2. Performance comparison of the LSTM, Transformer, and TPC architectures"
        + (f" when training on {_fmt_n(n_train)} patients from the BlendedICU dataset" if n_train else "")
    )
    footnote = (
        "AUROC: area under the receiver operating characteristic curve; MAPE: mean absolute percentage error; "
        "CI: confidence interval; TPC: temporal pointwise convolution; LSTM: long short-term memory. "
        "For each architecture, a single model was trained on an equal-n extract from all four databases; "
        "validation is reported separately by source database. "
        "AUROC 95% CI: DeLong (stay-level); MAPE 95% CI: patient-level bootstrap 10,000 (percentile). "
        "Best value per row in bold (lowest MAPE; highest AUROC). "
        f"Cohort: {args.cohort_label}."
        + (f" CI models available: {n_models_ci}/3." if n_models_ci and n_models_ci < 3 else "")
    )

    out = Path(args.out)
    if not str(out).lower().endswith(".xlsx"):
        out = Path(str(out) + ".xlsx")
    write_s2_xlsx(display, best_src, out, title, footnote)
    display.to_csv(out.with_name("Table_S2_moesm2.csv"), index=False)
    print("wrote", out)
    print(display.to_string(index=False))
    if n_models_ci < 3:
        print(f"WARN: CI file has {n_models_ci}/3 models — re-run eval_liver_paper_ci on equaln", file=sys.stderr)


if __name__ == "__main__":
    main()
